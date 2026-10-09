"""Read exact XLSX native cells; emit source-local tables without modifying originals."""
import argparse
from collections import Counter, defaultdict
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/source-candidates/svara-contacts-20261002"
OUT = ROOT / "public/data/svara_synaptic_contacts_v1"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHA = lambda data: hashlib.sha256(data).hexdigest()

def read_source(n):
    name = f"1-s2.0-S2211124718307563-mmc{n}.xlsx"
    data = (SOURCE / name).read_bytes()
    receipt = json.loads((SOURCE / (name + ".receipt.json")).read_bytes())
    assert SHA(data) == receipt["sha256"] and len(data) == receipt["bytes"]
    assert receipt["publishedSizeAndMd5Verified"] and receipt["status"] == 200
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert z.testzip() is None
        assert not any("externalLink" in p or "vbaProject" in p or "embedding" in p for p in z.namelist())
        xml = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        assert not xml.findall(".//s:f", NS), "Native source contains an unexpected formula"
        tokens = {c.attrib["r"]: c.find("s:v", NS).text for c in xml.findall(".//s:c", NS)
                  if c.attrib.get("t", "n") == "n" and c.find("s:v", NS) is not None}
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    assert len(wb.sheetnames) == 1
    sheet = wb.active
    values = list(sheet.values)
    info = {"componentId": receipt["componentId"], "name": name, "sheet": sheet.title,
            "range": sheet.calculate_dimension(), "bytes": len(data), "sha256": SHA(data),
            "md5": receipt["md5"], "sourceUrl": receipt["sourceUrl"]}
    return values, tokens, info

def build():
    old = json.loads((ROOT / "data/svara-supplement-catalogues-0.1.json").read_bytes())
    metadata = json.loads((ROOT / "data/source-candidates/svara-supplements-20261002/item-metadata.json").read_bytes())
    assert metadata["objectId"] == "item_3055210" and metadata["versionNumber"] == 7 and metadata["versionState"] == "RELEASED"
    assert next(x for x in metadata["files"] if x["objectId"] == "file_3055224")["metadata"]["license"] == "http://creativecommons.org/licenses/by/4.0/"
    policy_dir = ROOT / "data/source-candidates/svara-supplements-20261002"
    policy = json.loads((policy_dir / "publisher-supplement-policy.receipt.json").read_bytes())
    page = (policy_dir / "publisher-supplement-policy.html").read_bytes()
    assert policy["status"] == 200 and SHA(page) == policy["sha256"] and policy["verifiedPhrase"].encode() in page
    doc = {"schemaVersion": "0.1.0", "contractId": "svara-synaptic-contacts-0.1",
           "sourceId": old["sourceId"], "releaseId": old["releaseId"], "citation": old["citation"],
           "biologicalContext": old["biologicalContext"],
           "representation": {"kind": "published_source_local_numeric_tables", "coordinatesAvailable": False,
                              "editableGeometry": False, "physicalCellMappingEstablished": False,
                              "functionalTransmissionEstablished": False, "sameKohAnnotationVersionEstablished": False},
           "rights": {"license": "CC-BY-4.0", "licenseUrl": old["rights"]["licenseUrl"],
                      "basis": "Exact official article license plus explicit publisher supplementary-data policy applied to these four byte-verified public MPG v7 components.",
                      "articleLicenseLocator": old["rights"]["articleLicenseLocator"],
                      "publisherPolicyUrl": policy["sourceUrl"], "publisherPolicySha256": policy["sha256"],
                      "metadataSha256": SHA((policy_dir / "item-metadata.json").read_bytes()),
                      "separateComponentLicenseFieldPresent": False, "freidokWrapperCovered": False,
                      "kohAnnotationVersionCovered": False,
                      "modifications": "Native labels and numeric OOXML tokens extracted into JSON; counts and sums derived. Source workbooks unchanged."},
           "interpretation": {"contactAreaUnit": "µm²", "unitLocator": {
                                  "CiD": "file_3055227 / CiD to MN Contacts!C1: Synaptic Contact Area [µm²]",
                                  "CoBL": "file_3055229 / CoBL to MN Contacts!C1: Synaptic Contact Area [µm²]"},
                              "matrixUnitExplicitInWorkbook": False,
                              "matrixMeaning": "Published wiring values numerically agree with summed native contact areas; counts are derived separately from contact rows.",
                              "agreementTolerance": "0.000000001", "tolerancePurpose": "Source decimal reconciliation, not a biological threshold.",
                              "zeroMeaning": "Literal matrix zero and no listed contacts for this source pair; not proof of biological absence.",
                              "rowIdentity": "Exact component + sheet + native row; repeated endpoint pairs are distinct contacts, not duplicates."},
           "groups": []}
    outputs = {}
    for group, matrix_n, contacts_n in [("CiD", 2, 3), ("CoBL", 4, 5)]:
        matrix, matrix_tokens, matrix_info = read_source(matrix_n)
        contacts, contact_tokens, contacts_info = read_source(contacts_n)
        for info in [matrix_info, contacts_info]:
            native = next(x for x in metadata["files"] if x["objectId"] == info["componentId"])
            assert native["name"] == info["name"] and native["size"] == info["bytes"] and native["checksum"] == info["md5"]
        pre = [row[0] for row in matrix[1:]]
        post = list(matrix[0][1:])
        assert len(set(pre)) == len(pre) and len(set(post)) == len(post) and len(post) == 68
        assert contacts[0] == ("Presynaptic Cell", "Postsynaptic Cell", "Synaptic Contact Area [µm²]")
        values = [[matrix_tokens[f"{openpyxl.utils.get_column_letter(j+2)}{i+2}"] for j in range(len(post))] for i in range(len(pre))]
        rows = [{"sourceRow": i+2, "pre": row[0], "post": row[1], "areaText": contact_tokens[f"C{i+2}"]} for i, row in enumerate(contacts[1:])]
        assert all(row["pre"] in pre and row["post"] in post and Decimal(row["areaText"]) > 0 for row in rows)
        sums = defaultdict(Decimal)
        counts = Counter()
        for row in rows:
            pair = (row["pre"], row["post"])
            sums[pair] += Decimal(row["areaText"])
            counts[pair] += 1
        errors = [abs(Decimal(values[i][j]) - sums[(p, q)]) for i, p in enumerate(pre) for j, q in enumerate(post)]
        assert max(errors) <= Decimal(doc["interpretation"]["agreementTolerance"])
        assert all((Decimal(values[i][j]) == 0) == (counts[(p,q)] == 0) for i,p in enumerate(pre) for j,q in enumerate(post))
        summary = {"presynapticEntries": len(pre), "postsynapticEntries": len(post), "matrixCells": len(pre)*len(post),
                   "literalZeros": sum(Decimal(v) == 0 for row in values for v in row), "missingMatrixCells": 0,
                   "contactRows": len(rows), "distinctListedPairs": len(counts),
                   "maxMatrixAreaSumAbsoluteDifference": str(max(errors)), "pairsCompared": len(errors)}
        payload = {"schemaVersion": doc["schemaVersion"], "sourceId": doc["sourceId"], "releaseId": doc["releaseId"], "group": group,
                   "matrix": {**matrix_info, "pre": pre, "post": post, "values": values},
                   "contacts": {**contacts_info, "headers": list(contacts[0]), "rows": rows}, "summary": summary}
        data = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode()
        doc["groups"].append({"id": group, "pre": pre, "post": post, "matrixSource": matrix_info, "contactsSource": contacts_info,
                              "summary": summary, "asset": {"url": f"/data/svara_synaptic_contacts_v1/{group}.json", "bytes": len(data), "sha256": SHA(data)}})
        outputs[OUT / f"{group}.json"] = data
    outputs[ROOT / "data/svara-synaptic-contacts-0.1.json"] = (json.dumps(doc, ensure_ascii=False, indent=2) + "\n").encode()
    return outputs, doc

if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--write", action="store_true"); args = parser.parse_args()
    outputs, doc = build()
    for path, data in outputs.items():
        if args.write:
            path.parent.mkdir(exist_ok=True, parents=True); path.write_bytes(data)
        else:
            assert path.read_bytes() == data, f"Derivative drift: {path}"
    print(json.dumps({"mode": "write" if args.write else "check", "groups": [{"id": x["id"], **x["summary"], "assetBytes": x["asset"]["bytes"]} for x in doc["groups"]]}))
